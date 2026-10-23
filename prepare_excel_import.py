"""Read original workbook, prepare provenance-preserving historic shop records. No XLSX writes."""
import collections, datetime as dt, decimal, glob, hashlib, json, math, pathlib, re, uuid, sys
sys.stdout.reconfigure(encoding='utf-8')

ROOT=pathlib.Path(__file__).parent
OUT=ROOT/'imports'/'excel-all-years'
OUT.mkdir(parents=True,exist_ok=True)
import openpyxl
source=next(ROOT.glob('*.xlsx'))
formulas=openpyxl.load_workbook(source,data_only=False)
values=openpyxl.load_workbook(source,data_only=True)
issues=[]; skips=[]; sections=[]; headers={}; corrections=collections.Counter(); source_controls=collections.defaultdict(lambda:dict(rows=0,revenue=0,cost=0,expense=0,profit=0))
products=['จิ๋ว','กลม','อั่งท้อ','แดง','ตัด']
uid=lambda: str(uuid.uuid4())
def num(value):
    if value is None or value=='':return 0.0
    if isinstance(value,dt.datetime):return float(openpyxl.utils.datetime.to_excel(value))
    if isinstance(value,(int,float)) and not isinstance(value,bool) and math.isfinite(value):return float(value)
    if isinstance(value,str) and re.fullmatch(r'-?\d+(?:,\d{3})*(?:\.\d+)?',value.strip()):return float(value.replace(',',''))
    raise ValueError(f'Non-numeric value: {value}')
def day(value,year):
    if isinstance(value,(dt.datetime,dt.date)):
        y=value.year
        if y>=2400:y-=543
        elif 1900<=y<2000 and y%100==(year+543)%100:y=year
        result=dt.date(y,value.month,value.day)
        return result.isoformat()
    if isinstance(value,str):
        m=re.match(r'\s*(\d{1,2})(มค|กพ|มีค|เมย|พค|มิย|กค|สค|กย|ตค|พย|ธค)(\d{2,4})',value.replace('.','').replace(' ',''))
        if m:
            months=['มค','กพ','มีค','เมย','พค','มิย','กค','สค','กย','ตค','พย','ธค'];y=int(m[3]);y=2500+y if y<100 else y;y=y-543 if y>2400 else y
            return dt.date(y,months.index(m[2])+1,int(m[1])).isoformat()
    return None
location_names={}
def canonical(name):
    text=re.sub(r'\s+',' ',str(name).strip())
    key=re.sub(r'\s+','',text).casefold()
    aliases={'centralbangna':'Central Bangna','centralrama2':'Central Rama2','centralrama3':'Central Rama3','centraleastville':'Central Eastville','centralwestgate':'Central Westgate','centralladprao':'Central Ladprao','centralchidlom':'Central Chidlom','centralpinklao':'Central Pinklao','centralchaengwattana':'Central Chaengwattana','centralwestville':'Central Westville','seaconbangkhae':'Seacon Bangkae','seaconbangkae':'Seacon Bangkae','themallthaphra':'The Mall ท่าพระ','themallthapra':'The Mall ท่าพระ','themallท่าพระ':'The Mall ท่าพระ','themallbangkae':'The Mall บางแค','themallngam':'The Mall งามวงศ์วาน','themallงาม':'The Mall งามวงศ์วาน'}
    text=aliases.get(key,text)
    key=re.sub(r"[\s’']+",'',text).casefold()
    return location_names.setdefault(key,text)
def numeric_payload(ws,r,start,end):
    return any(isinstance(ws.cell(r,c).value,(int,float)) and ws.cell(r,c).value!=0 for c in range(start,end+1))
def rates(formula,row,cols):
    if not isinstance(formula,str) or not formula.startswith('='):return {}
    result={}
    for name,col in zip(products,cols):
        ref=rf'\$?{openpyxl.utils.get_column_letter(col)}\$?{row}(?!\d)'
        m=re.search(ref+ r'\s*\*\s*(\d+(?:\.\d+)?)',formula,re.I)
        if not m:m=re.search(r'(\d+(?:\.\d+)?)\s*\*\s*'+ref,formula,re.I)
        if m:result[name]=float(m[1])
    return result

latest_cost={};latest_price={};seen_fingerprints={};all_rows=[]
for sf in formulas:
    if not sf.title.startswith('Event'):continue
    sv=values[sf.title];year=int(sf.title[-4:])-543;current=None;shift=0;active_headers={}
    for r in range(1,sf.max_row+1):
        header_col=next((c for c in [3,4] if str(sf.cell(r,c).value).strip()=='จิ๋ว' and str(sf.cell(r,c+1).value).strip()=='กลม'),None)
        if header_col:
            shift=header_col-3
            label=sf.cell(r,1+shift).value
            if label and not isinstance(label,dt.datetime):
                current=dict(id=uid(),name=canonical(label),sheet=sf.title,header_row=r,rows=[])
                sections.append(current)
            active_headers={c:str(sf.cell(r,c).value).strip() for c in range(8+shift,22+shift) if sf.cell(r,c).value is not None}
            continue
        # Some historic wholesale sections only have a location label, without a column header.
        label=sf.cell(r,1+shift).value
        if isinstance(label,str) and label.strip() and not numeric_payload(sv,r,3+shift,22+shift):
            next_date=day(sf.cell(r+1,2+shift).value,year) if r<sf.max_row else None
            if next_date and not re.search(r'^(mon|tue|wed|thu|fri|sat|sun|รวม|สรุป)',label.strip(),re.I):
                current=dict(id=uid(),name=canonical(label),sheet=sf.title,header_row=r,rows=[]);sections.append(current)
        raw_date=sf.cell(r,2+shift).value
        date=day(raw_date,year)
        if not date:
            if isinstance(raw_date,str) and raw_date.strip() and numeric_payload(sv,r,3+shift,22+shift):
                skips.append(dict(sheet=sf.title,row=r,reason='unparseable date',raw=str(raw_date)))
            continue
        if not numeric_payload(sv,r,3+shift,22+shift):
            skips.append(dict(sheet=sf.title,row=r,reason='empty / zero template row'));continue
        if isinstance(raw_date,dt.datetime) and raw_date.year!=int(date[:4]):corrections['date_year_normalized']+=1
        elif isinstance(raw_date,str):corrections['thai_text_dates']+=1
        if current is None:raise ValueError(f'{sf.title}:{r} has no location')
        def n(c):
            cell=sv.cell(r,c+shift)
            if isinstance(cell.value,dt.datetime):
                ref=f'{sf.title}!{cell.coordinate}'
                if ref not in headers:headers[ref]=True;issues.append(dict(sheet=sf.title,row=r,type='currency_cell_formatted_as_date',cell=cell.coordinate,value=num(cell.value)))
            return num(cell.value)
        revcol,expcol,profitcol=(17,16,18) if year==2020 else (13,14,22)
        # A March 2021 wholesale block reverses Revenue and Expense in the original formulas.
        profit_formula=str(sf.cell(r,15+shift).value)
        if year==2021 and re.search(r'\bN'+str(r)+r'\s*-\s*M'+str(r)+r'\b',profit_formula,re.I):
            revcol,expcol=14,13;corrections['reversed_revenue_expense_columns']+=1
        revenue=n(revcol);raw_cost=n(8)
        if revenue<0 or raw_cost<0:raise ValueError(f'Negative revenue/cost {sf.title}:{r}')
        cols=[c+shift for c in range(3,8)];cost_rates=rates(sf.cell(r,8+shift).value,r,cols)
        quantities={name:n(c) for name,c in zip(products,range(3,8))}
        if any(q<0 for q in quantities.values()):raise ValueError(f'Negative quantity {sf.title}:{r}')
        priced_cost=sum(quantities[name]*cost_rates.get(name,0) for name in products)
        items=[]
        if abs(priced_cost-raw_cost)<0.005:
            items=[dict(name=name,quantity=q,cost=cost_rates.get(name,0),price=0) for name,q in quantities.items() if q]
        else:
            items=[dict(name=name,quantity=q,cost=0,price=0) for name,q in quantities.items() if q]
            if raw_cost:items.append(dict(name='ต้นทุนวัตถุดิบรวมจาก Excel',quantity=1,cost=raw_cost,price=0))
            if raw_cost:issues.append(dict(sheet=sf.title,row=r,type='aggregate_cost',message='Unit costs cannot be verified; source total cost retained separately'))
        expense_cols=list(range(9,16)) if year==2020 else list(range(9,13))+list(range(16,22))
        expenses=[]
        for c in expense_cols:
            amount=n(c)
            if amount<0:raise ValueError(f'Negative expense {sf.title}:{r}:{c}')
            if amount:expenses.append(dict(name=active_headers.get(c+shift,f'รายจ่าย {openpyxl.utils.get_column_letter(c+shift)}'),amount=amount))
        base_categories=raw_cost+sum(n(c) for c in (range(9,16) if year==2020 else range(9,13)))
        source_expense=sv.cell(r,expcol+shift).value
        if source_expense is not None:
            difference=n(expcol)-base_categories
            if abs(difference)>.005:
                issues.append(dict(sheet=sf.title,row=r,type='expense_total_mismatch',source=n(expcol),categories=base_categories))
        cost=sum(i['quantity']*i['cost'] for i in items);expense=sum(e['amount'] for e in expenses);profit=revenue-cost-expense
        source_profit=sv.cell(r,profitcol+shift).value
        if source_profit is not None and abs(num(source_profit)-profit)>.011:
            issues.append(dict(sheet=sf.title,row=r,type='profit_formula_mismatch',source=num(source_profit),calculated=round(profit,2)))
        for name,rate in cost_rates.items():
            if date>=latest_cost.get(name,('',0))[0]:latest_cost[name]=(date,rate)
        # Do not infer retail selling prices from wholesale revenues or total daily sales.
        fingerprint=json.dumps([current['name'].casefold(),date,quantities,round(revenue,4),round(cost,4),round(expense,4)],ensure_ascii=False,sort_keys=True)
        if fingerprint in seen_fingerprints:
            skips.append(dict(sheet=sf.title,row=r,reason='exact duplicate historic entry',duplicate_of=seen_fingerprints[fingerprint]));continue
        seen_fingerprints[fingerprint]=f'{sf.title}!{r}'
        row_issues=[i for i in issues if i['sheet']==sf.title and i['row']==r]
        note=f'นำเข้าจาก {source.name} | {sf.title}!แถว {r} | ข้อมูลย้อนหลัง: ไม่มีข้อมูลของเหลือ ไม่ตัดสต็อกอัตโนมัติ'
        if str(raw_date)[:10]!=date:note+=' | วันที่ต้นฉบับ: '+str(raw_date)
        if source_profit is not None:note+=' | กำไรเดิมใน Excel: '+str(source_profit)
        if row_issues:note+=' | ตรวจสอบ: '+', '.join(sorted(set(i['type'] for i in row_issues)))
        record=dict(id=uid(),period_id=current['id'],sale_date=date,revenue=revenue,items=items,expenses=expenses,note=note)
        current['rows'].append(record);all_rows.append(record)
        control=source_controls[sf.title];control['rows']+=1
        for key,amount in [('revenue',revenue),('cost',cost),('expense',expense),('profit',profit)]:control[key]+=amount

# Split duplicate dates within the same section into separate explicit periods, never overwrite.
periods=[];locations={};records=[]
for section in sections:
    if not section['rows']:continue
    location=locations.setdefault(section['name'],dict(id=uid(),name=section['name']))
    groups=[]
    for row in section['rows']:
        group=next((g for g in groups if row['sale_date'] not in {r['sale_date'] for r in g}),None)
        if group is None:group=[];groups.append(group)
        group.append(row)
    for index,group in enumerate(groups):
        pid=section['id'] if index==0 else uid();dates=[r['sale_date'] for r in group]
        periods.append(dict(id=pid,location_id=location['id'],name=f"นำเข้า {section['sheet']} แถว {section['header_row']}"+(f' ชุด {index+1}' if len(groups)>1 else ''),starts_on=min(dates),ends_on=max(dates),daily_rent=0,daily_wage=0))
        for row in group:row['period_id']=pid;records.append(row)
        if index:issues.append(dict(sheet=section['sheet'],row=section['header_row'],type='duplicate_dates_split',message='Separate period preserves multiple source rows on same date'))

stock_items=[dict(id=uid(),name=name,unit='ชิ้น' if name!='ตัด' else 'ถาด',kind='product',cost=latest_cost.get(name,('',0))[1],price=0) for name in products]
product_ids={i['name']:i['id'] for i in stock_items}
for record in records:
    for item in record['items']:
        if item['name'] in product_ids:item['product_id']=product_ids[item['name']]
stock=values['Stock'];movements=[];stock_location_aliases={'บางนา':'Central Bangna','พระราม2':'Central Rama2','อีสวิว':'Central Eastville','พระราม3':'Central Rama3','ท่าพระ':'The Mall ท่าพระ','บางแค':'Seacon Bangkae','ลาดพร้าว':'Central Ladprao','ทรู':'ทรู','new':'new'}
for r in range(2,33):
    name=str(stock.cell(r,1).value or '').strip()
    if not name:continue
    item=dict(id=uid(),name=name,unit='ชิ้น',kind='equipment',cost=0,price=0);stock_items.append(item);allocated=0
    for c in range(3,12):
        quantity=num(stock.cell(r,c).value)
        if not quantity:continue
        label=canonical(stock_location_aliases[str(stock.cell(1,c).value)])
        location=locations.setdefault(label,dict(id=uid(),name=label));allocated+=quantity
        movements.append(dict(id=uid(),stock_item_id=item['id'],location_id=location['id'],moved_on='2026-09-10',quantity=quantity,reason=f'ยอดอุปกรณ์ตาม Excel Stock!{openpyxl.utils.get_column_letter(c)}{r} อัปเดต 10/9/26'))
    total=stock.cell(r,2).value
    if isinstance(total,str) and re.fullmatch(r'\d+(?:\+\d+)+',total):total=sum(map(int,total.split('+')))
    if total is not None:
        difference=num(total)-allocated
        if difference>0:
            location=locations.setdefault('อุปกรณ์ยังไม่ระบุสถานที่',dict(id=uid(),name='อุปกรณ์ยังไม่ระบุสถานที่'))
            movements.append(dict(id=uid(),stock_item_id=item['id'],location_id=location['id'],moved_on='2026-09-10',quantity=difference,reason=f'ส่วนต่าง Total กับจำนวนแยกสถานที่จาก Stock!B{r} ยังไม่ทราบสถานที่จริง'))
        elif difference<0:issues.append(dict(sheet='Stock',row=r,type='stock_allocation_exceeds_total',total=num(total),allocated=allocated))

plan=dict(source=source.name,source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),locations=list(locations.values()),periods=periods,daily_records=records,stock_items=stock_items,stock_movements=movements)
for table in ['locations','periods','daily_records','stock_items','stock_movements']:
    ids=[r['id'] for r in plan[table]];assert len(ids)==len(set(ids)),table
assert len(records)==sum(c['rows'] for c in source_controls.values())
period_map={p['id']:p for p in periods}
assert all(period_map[r['period_id']]['starts_on']<=r['sale_date']<=period_map[r['period_id']]['ends_on'] for r in records)
summary=dict(source=source.name,counts={t:len(plan[t]) for t in ['locations','periods','daily_records','stock_items','stock_movements']},controls=source_controls,corrections=corrections,issue_counts=collections.Counter(i['type'] for i in issues),skipped_counts=collections.Counter(i['reason'] for i in skips),excluded_sheets=[s for s in formulas.sheetnames if not s.startswith('Event') and s!='Stock'])
for name,data in [('plan.json',plan),('summary.json',summary),('issues.json',issues),('skipped.json',skips)]:
    (OUT/name).write_text(json.dumps(data,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
print(json.dumps(summary,ensure_ascii=False,indent=2))
